"""Integration tests for the export facade (reference-function parity).

The ``completed_run`` fixture lives in conftest.py and is shared with the
analysis-summary tests.
"""

from pathlib import Path

import pytest

from suite.db import RunService, make_engine, make_session_factory, run_migrations
from suite.db.payloads import read_payload
from suite.exports import run_exports_sync
from suite.schemas import RunDetailResponse


def test_export_csv_group(completed_run):
    svc, run_id, run_dir = completed_run
    artifacts = run_exports_sync(run_id, svc, ["csv"])
    csv_dir = Path(artifacts["csv"][0])
    assert (csv_dir / "benchmark_summary.csv").exists()
    assert (csv_dir / "benchmark_raw_costs.csv").exists()
    assert (csv_dir / "qa_friedman_objective_error.csv").exists()
    assert (csv_dir / "qe_scaling_analysis.csv").exists()
    # per-run CSVs reconstructed from payloads
    per_run = list(csv_dir.glob("**/run_*.csv"))
    assert len(per_run) == 8  # 2 runs x 4 algos


def test_export_charts_group(completed_run):
    svc, run_id, run_dir = completed_run
    artifacts = run_exports_sync(run_id, svc, ["charts"])
    chart_dir = Path(artifacts["charts"][0])
    pngs = list(chart_dir.glob("*.png"))
    assert len(pngs) > 0


def test_export_docx_group(completed_run):
    svc, run_id, run_dir = completed_run
    artifacts = run_exports_sync(run_id, svc, ["docx"])
    report = Path(artifacts["docx"][0])
    assert report.name == "benchmark_report.docx"
    assert report.stat().st_size > 10_000


def test_export_json_group(completed_run):
    svc, run_id, run_dir = completed_run
    artifacts = run_exports_sync(run_id, svc, ["json"])
    assert (Path(artifacts["json"][0])).exists()


def test_module_constants_restored(completed_run):
    """Every reference-module constant the export pipeline patches is
    restored after export (directories AND the run-configuration mirror)."""
    import hyde_bench.run_benchmark as rb

    snapshot = (
        rb.CSV_DIR,
        rb.CHART_DIR,
        rb.HERE,
        rb.N_RUNS,
        rb.MAX_EVALS,
        rb.ALPHA,
        rb.TEST_CASES,
    )
    svc, run_id, _ = completed_run
    run_exports_sync(run_id, svc, ["csv", "charts", "docx", "json"])
    assert (
        rb.CSV_DIR,
        rb.CHART_DIR,
        rb.HERE,
        rb.N_RUNS,
        rb.MAX_EVALS,
        rb.ALPHA,
        rb.TEST_CASES,
    ) == snapshot


def test_export_mirrors_run_configuration(completed_run):
    """Charts/DOCX describe the actual experiment: for a run configured with
    n_runs=2 the report's conditions and convergence counts must reflect 2
    runs, not the CLI default (50) the reference module hardcodes."""
    import docx

    svc, run_id, run_dir = completed_run
    run_exports_sync(run_id, svc, ["docx"])
    doc = docx.Document(str(run_dir / "benchmark_report.docx"))
    paragraphs = [p.text for p in doc.paragraphs]
    assert any("2 independent runs" in t for t in paragraphs)
    cell_texts = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    assert any("/2" in t for t in cell_texts), (
        "convergence counts do not use the run's n_runs=2 denominator"
    )


def test_docx_only_export_includes_charts(completed_run):
    """Selecting only the docx group still renders the charts the report
    embeds, and does not list them under the charts group."""
    svc, run_id, run_dir = completed_run
    artifacts = run_exports_sync(run_id, svc, ["docx"])
    assert set(artifacts.keys()) == {"docx"}
    report = Path(artifacts["docx"][0])
    assert report.exists() and report.stat().st_size > 0
    charts = run_dir / "benchmark_charts"
    assert any(charts.glob("*.png")), "report was produced without figures"


def test_json_artifacts_not_leaked_into_other_groups(completed_run):
    """Requesting csv only must not list json artifacts (analysis_summary.json
    is still written as the analysis cache, but not reported)."""
    svc, run_id, run_dir = completed_run
    artifacts = run_exports_sync(run_id, svc, ["csv"])
    assert set(artifacts.keys()) == {"csv"}
    assert (run_dir / "analysis_summary.json").exists()


def test_render_guard(tmp_path: Path):
    """When matplotlib's renderer geometry is inflated ("Image size ...
    too large" / "raster overflow" / bad_alloc), the export guard skips
    the overflowing tight_layout and saves via the dpi-reset + plain-bbox
    fallback instead of losing the chart."""
    import matplotlib.figure as mfig

    calls: list[tuple[float, object]] = []
    logged: list[str] = []
    original_savefig = mfig.Figure.savefig
    original_tight = mfig.Figure.tight_layout

    def fake_savefig(self, fname, *args, **kwargs):  # type: ignore[no-untyped-def]
        calls.append((self.dpi, kwargs.get("bbox_inches")))
        if self.dpi != 100:
            raise RuntimeError(
                "FT_Render_Glyph (ft2font_wrapper.cpp line 1943) failed "
                "with error 0x62: raster overflow"
            )
        if kwargs.get("bbox_inches") == "tight":
            raise ValueError(
                "Image size of 999x999 pixels is too large. "
                "It must be less than 2^23 in each direction."
            )
        return "saved"

    def fake_tight_layout(self, *args, **kwargs):  # type: ignore[no-untyped-def]
        raise MemoryError("std::bad_alloc")

    mfig.Figure.savefig = fake_savefig  # type: ignore[method-assign]
    mfig.Figure.tight_layout = fake_tight_layout  # type: ignore[method-assign]
    try:
        from suite.exports import _install_render_guard

        _install_render_guard(logged.append)
        fig = mfig.Figure()
        fig.dpi = 150.0  # emulate the inflated geometry observed in-app
        # tight_layout overflow is skipped, not propagated
        assert fig.tight_layout() is None
        result = fig.savefig("x.png", dpi=150, bbox_inches="tight")
        assert result == "saved"
        # first attempt kept the (corrupt) dpi, the recovery reset it to
        # 100 and dropped the tight bbox
        assert calls[0] == (150.0, "tight")
        assert calls[-1] == (100, None)
        assert any("tight bbox overflow" in m for m in logged)
        assert any("tight_layout skipped" in m for m in logged)
    finally:
        mfig.Figure.savefig = original_savefig  # type: ignore[method-assign]
        mfig.Figure.tight_layout = original_tight  # type: ignore[method-assign]


def _patched_make_charts(monkeypatch, exc: Exception) -> None:
    """Force the in-process ``make_charts`` step to always fail; the
    monkeypatch is process-local, so the clean-subprocess fallback still
    runs the real reference function."""
    import hyde_bench.run_benchmark as rb

    def broken(*args, **kwargs):
        raise exc

    monkeypatch.setattr(rb, "make_charts", broken)


def test_chart_subprocess_fallback(completed_run, monkeypatch):
    """When the renderer-geometry corruption makes even the in-process
    plain-bbox fallback fail, the chart step is re-run in a clean
    subprocess and the export still succeeds with every PNG present."""
    _patched_make_charts(
        monkeypatch,
        RuntimeError(
            "FT_Render_Glyph (ft2font_wrapper.cpp line 1943) failed "
            "with error 0x62: raster overflow"
        ),
    )
    svc, run_id, run_dir = completed_run
    artifacts = run_exports_sync(run_id, svc, ["charts"])
    chart_dir = Path(artifacts["charts"][0])
    pngs = {p.name for p in chart_dir.glob("*.png")}
    # make_charts is the patched step; the subprocess fallback must have
    # produced its per-benchmark and summary charts
    assert "booth_2D.png" in pngs
    assert "qd_margin_wins.png" in pngs
    assert "qa_objective_error_wins.png" in pngs
    # the other chart steps rendered in-process as usual
    assert "qc_wall_time_per_benchmark.png" in pngs
    log = (run_dir / "export_error.log").read_text(encoding="utf-8")
    assert "clean subprocess" in log


def test_chart_subprocess_not_triggered_for_other_errors(completed_run, monkeypatch):
    """A chart step that fails with a non-render error must not waste a
    clean-subprocess retry; the failure is recorded and the export surfaces
    it to the UI instead of pretending success."""
    _patched_make_charts(monkeypatch, RuntimeError("boom"))
    svc, run_id, run_dir = completed_run
    with pytest.raises(RuntimeError, match="make_charts"):
        run_exports_sync(run_id, svc, ["charts"])
    log = (run_dir / "export_error.log").read_text(encoding="utf-8")
    assert "boom" in log
    assert "clean subprocess" not in log


def test_export_rejected_while_run_in_progress(tmp_path, session_factory):
    """A run whose worker is still active must not be exported: its
    benchmark_results.json and payloads are still being written."""
    from suite.db import RunService

    svc = RunService(session_factory)
    run_dir = tmp_path / "runs" / "in_progress"
    run_dir.mkdir(parents=True)
    run = svc.create_run(
        label="in progress",
        output_dir=str(run_dir),
        n_runs=2,
        max_evals=1000,
        alpha=0.05,
        seed_base=0,
        algo_params={},
        test_cases=[{"fname": "booth", "dim": 2}],
    )
    svc.mark_running(run.id)
    with pytest.raises(RuntimeError, match="still in progress"):
        run_exports_sync(run.id, svc, ["csv"])


def test_payload_shape_matches_frontend_schema(completed_run):
    """The persisted payload must contain every field the frontend zod
    scenarioPayloadSchema requires (drift broke charts before)."""

    svc, run_id, run_dir = completed_run
    detail = RunDetailResponse(**svc.get_run_detail(run_id))
    required = {
        "raw_costs",
        "raw_wall_ms",
        "raw_evals",
        "raw_aucs",
        "raw_obj_errors",
        "mean_curve",
        "curves",
        "conv_binary",
        "global_opt",
    }
    for sr in detail.scenario_results:
        payload = read_payload(Path(sr.payloads_path))
        missing = required - set(payload.keys())
        assert not missing, f"{sr.algo_key} payload missing: {missing}"
        assert isinstance(payload.get("replay_histories", []), list)


def test_export_missing_results_file(tmp_path: Path):
    db_path = tmp_path / "suite.db"
    engine = make_engine(db_path)
    run_migrations(engine, db_path)
    svc = RunService(make_session_factory(engine))
    run_dir = tmp_path / "runs" / "empty"
    run_dir.mkdir(parents=True)
    run = svc.create_run(
        label="empty",
        output_dir=str(run_dir),
        n_runs=1,
        max_evals=1000,
        alpha=0.05,
        seed_base=0,
        algo_params={},
        test_cases=[{"fname": "booth", "dim": 2}],
    )
    svc.mark_running(run.id)
    svc.mark_completed(run.id, 1.0)
    with pytest.raises(FileNotFoundError):
        run_exports_sync(run.id, svc, ["csv"])
    engine.dispose()
